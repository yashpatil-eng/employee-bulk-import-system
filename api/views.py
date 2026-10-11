
import csv
import io
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from django.conf import settings
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import HttpResponse

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from rest_framework import status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response

from .models import Employee
from .serializers import RegisterSerializer


BATCH_SIZE = 500
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def read_rows(uploaded_file, filename):
    """Yield spreadsheet rows without loading the whole workbook into memory."""

    if filename.endswith(".csv"):
        text_file = io.TextIOWrapper(
            uploaded_file.file,
            encoding="utf-8-sig",
            newline="",
        )
        try:
            yield from csv.reader(text_file)
        finally:
            # Keep the underlying uploaded file open for Django.
            try:
                text_file.detach()
            except ValueError:
                pass
        return

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )
    try:
        worksheet = workbook.active
        yield from worksheet.iter_rows(values_only=True)
    finally:
        workbook.close()


def text_value(value):
    return str(value).strip() if value is not None else ""


def validate_row(row, seen_emails):
    """Return cleaned employee data and a list of validation errors."""

    name = text_value(row.get("name"))
    email = text_value(row.get("email")).lower()
    department = text_value(row.get("department")) or None
    salary_text = text_value(row.get("salary"))

    errors = []
    salary = None

    if not name:
        errors.append("Name is required.")
    elif len(name) > 100:
        errors.append("Name cannot exceed 100 characters.")

    if not email:
        errors.append("Email is required.")
    else:
        try:
            validate_email(email)
        except ValidationError:
            errors.append("Email format is invalid.")

        if len(email) > 254:
            errors.append("Email cannot exceed 254 characters.")

    if department and len(department) > 100:
        errors.append("Department cannot exceed 100 characters.")

    if salary_text:
        try:
            salary = Decimal(salary_text)

            if not salary.is_finite() or salary < 0:
                errors.append("Salary must be a non-negative number.")
                salary = None
            elif salary.as_tuple().exponent < -2:
                errors.append("Salary cannot have more than 2 decimal places.")
                salary = None
            elif salary >= Decimal("10000000000"):
                errors.append("Salary exceeds the maximum supported value.")
                salary = None

        except (InvalidOperation, ValueError):
            errors.append("Salary must be a valid number.")

    if email:
        if email in seen_emails:
            errors.append("Duplicate email in uploaded file.")
        else:
            seen_emails.add(email)

    return {
        "name": name,
        "email": email,
        "department": department,
        "salary": salary,
    }, errors


def database_duplicate_emails(emails):
    """Find existing emails without making one query per row."""

    if not emails:
        return set()

    # Emails are normalized to lowercase before import.
    # iexact also catches existing mixed-case email records.
    query = Q()
    for email in emails:
        query |= Q(email__iexact=email)

    return set(
        Employee.objects.filter(query).values_list("email", flat=True)
    )


def save_batch(valid_rows, results, imported_count):
    """Bulk insert a batch; use individual inserts to identify conflicts."""

    if not valid_rows:
        return imported_count

    try:
        with transaction.atomic():
            Employee.objects.bulk_create(
                [item["object"] for item in valid_rows]
            )

        for item in valid_rows:
            results[item["result_index"]]["imported"] = True
            imported_count += 1

    except IntegrityError:
        # Rollback the failed batch, then identify conflicting rows.
        for item in valid_rows:
            result = results[item["result_index"]]

            try:
                with transaction.atomic():
                    item["object"].save(force_insert=True)

                result["imported"] = True
                imported_count += 1

            except IntegrityError:
                result["errors"].append(
                    "Email already exists in the database."
                )
                result["is_valid"] = False

    return imported_count


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def import_employees(request):
    uploaded_file = request.FILES.get("file")

    if not uploaded_file:
        return Response(
            {"error": "Please upload a file."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    filename = uploaded_file.name.lower()

    if not filename.endswith((".csv", ".xlsx")):
        return Response(
            {"error": "Only CSV and XLSX files are supported."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    if uploaded_file.size > MAX_FILE_SIZE:
        return Response(
            {"error": "File size must not exceed 10 MB."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        row_iterator = iter(read_rows(uploaded_file, filename))

        try:
            raw_headers = next(row_iterator)
        except StopIteration:
            return Response(
                {"error": "The uploaded file is empty."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        headers = [
            text_value(value).lower()
            for value in raw_headers
        ]

        if not any(headers):
            return Response(
                {"error": "The file has no column headers."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if len(headers) != len(set(headers)):
            return Response(
                {"error": "The file contains duplicate column headers."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        missing = {"name", "email"} - set(headers)
        if missing:
            return Response(
                {"error": "Missing required columns: " + ", ".join(sorted(missing))},
                status=status.HTTP_400_BAD_REQUEST,
            )

        results = []
        seen_emails = set()
        imported_count = 0
        batch = []
        row_number = 1

        for row_number, values in enumerate(row_iterator, start=2):
            values = list(values)

            # Ignore completely empty rows.
            if all(not text_value(value) for value in values):
                continue

            # Reject rows with missing or extra columns.
            if len(values) != len(headers):
                results.append({
                    "row": row_number,
                    "name": None,
                    "email": None,
                    "department": None,
                    "salary": None,
                    "errors": [
                        f"Expected {len(headers)} columns but found {len(values)}."
                    ],
                    "is_valid": False,
                    "imported": False,
                })
                continue

            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }

            cleaned, errors = validate_row(row, seen_emails)

            result = {
                "row": row_number,
                "name": cleaned["name"] or None,
                "email": cleaned["email"] or None,
                "department": cleaned["department"],
                "salary": (
                    str(cleaned["salary"])
                    if cleaned["salary"] is not None
                    else None
                ),
                "errors": errors,
                "is_valid": not errors,
                "imported": False,
            }

            results.append(result)

            if errors:
                continue

            batch.append({
                "object": Employee(
                    owner=request.user,
                    name=cleaned["name"],
                    email=cleaned["email"],
                    department=cleaned["department"],
                    salary=cleaned["salary"],
                ),
                "result_index": len(results) - 1,
            })

            if len(batch) >= BATCH_SIZE:
                emails = [item["object"].email for item in batch]
                existing = database_duplicate_emails(emails)

                remaining = []
                for item in batch:
                    email = item["object"].email
                    if any(saved.lower() == email for saved in existing):
                        result = results[item["result_index"]]
                        result["errors"].append(
                            "Email already exists in the database."
                        )
                        result["is_valid"] = False
                    else:
                        remaining.append(item)

                imported_count = save_batch(
                    remaining, results, imported_count
                )
                batch = []

        # Save the final partial batch.
        if batch:
            emails = [item["object"].email for item in batch]
            existing = database_duplicate_emails(emails)

            remaining = []
            for item in batch:
                email = item["object"].email
                if any(saved.lower() == email for saved in existing):
                    result = results[item["result_index"]]
                    result["errors"].append(
                        "Email already exists in the database."
                    )
                    result["is_valid"] = False
                else:
                    remaining.append(item)

            imported_count = save_batch(
                remaining, results, imported_count
            )

        failed_records = [
            item for item in results if item["errors"]
        ]

        return Response({
            "message": "File processing completed.",
            "filename": uploaded_file.name,
            "total_rows": len(results),
            "imported_rows": imported_count,
            "failed_rows": len(failed_records),
            "employees": results,
            "failed_records": failed_records,
        })

    except (UnicodeDecodeError, csv.Error):
        return Response(
            {"error": "Could not read the CSV file. Check its encoding."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    except (
        InvalidFileException,
        BadZipFile,
        KeyError,
        ValueError,
        OSError,
    ):
        return Response(
            {"error": "Could not read the file. Check its format."},
            status=status.HTTP_400_BAD_REQUEST,
        )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def download_failed_report(request):
    failed_rows = request.data.get("failed_records", [])

    if not isinstance(failed_rows, list):
        return Response(
            {"error": "failed_records must be a list."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    response = HttpResponse(
        content_type="text/csv; charset=utf-8"
    )
    response["Content-Disposition"] = (
        'attachment; filename="failed_employees.csv"'
    )
    response.write("\ufeff")

    writer = csv.writer(response)
    writer.writerow([
        "row", "name", "email", "department", "salary", "errors"
    ])

    for employee in failed_rows:
        if not isinstance(employee, dict):
            continue

        errors = employee.get("errors", [])
        if not isinstance(errors, list):
            errors = [str(errors)]

        writer.writerow([
            employee.get("row", ""),
            employee.get("name") or "",
            employee.get("email") or "",
            employee.get("department") or "",
            employee.get("salary") or "",
            "; ".join(str(error) for error in errors),
        ])

    return response

@api_view(["POST"])
@permission_classes([AllowAny])
def register_user(request):
    serializer = RegisterSerializer(data=request.data)

    if serializer.is_valid():
        user = serializer.save()

        return Response(
            {
                "message": "Registration successful.",
                "user": {
                    "id": user.id,
                    "username": user.username,
                    "email": user.email,
                },
            },
            status=status.HTTP_201_CREATED,
        )

    return Response(
        serializer.errors,
        status=status.HTTP_400_BAD_REQUEST,
    )